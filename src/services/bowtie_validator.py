"""Bow-tie validation: compare a scored cross-plane profile + HVL against the
measured RaySafe data.

The measured reference lives in the Oct 2023 workbook
(``research/2023 - CBCT Dose PCXMC/New results Oct 2023 (...).xlsx``), sheet
``Collated results``: lateral position (cm) in column C, and each CBCT mode
occupies a 4-column group (Dose uGy, uGy/s, HVL mmAl, blank) starting at column
E. The MC side is a CSV of (position_cm, dose) produced by an isocenter-plane
scorer. This service normalises both to peak=1, emits a comparison CSV, and
draws a matplotlib PNG.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass
from typing import Optional

import numpy as np

logger = logging.getLogger(__name__)


@dataclass
class Profile:
    """A cross-plane dose profile with optional per-position HVL."""

    label: str
    position_cm: np.ndarray
    dose: np.ndarray
    hvl_mmAl: Optional[np.ndarray] = None


def load_measured_profile(
    xlsx_path: str,
    mode_group: int = 0,
    sheet: str = "Collated results",
) -> Profile:
    """Load one mode's measured cross-plane profile from the RaySafe workbook.

    The ``Collated results`` sheet lays each CBCT mode out as a 4-column group
    (``Dose uGy``, ``uGy/s``, ``HVL (mm Al)``, blank), but the groups are NOT
    uniformly spaced -- Spotlight/Pelvis carry extra "Coast"/"Scaling" columns.
    This function locates the ``Dose uGy`` columns by scanning the header row
    and selects the ``mode_group``-th one (0 = Head, 1 = Spotlight, ...).

    Args:
        xlsx_path: Path to the ``New results Oct 2023 (...)`` workbook.
        mode_group: Index into the list of ``Dose uGy`` columns (0 = first mode).
        sheet: Sheet name (default ``Collated results``).
    """
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = wb[sheet]
    pos: list[float] = []
    dose: list[float] = []
    hvl: list[float] = []

    # First pass: find the header row and the Dose uGy columns. The groups are
    # unevenly spaced, so we cannot use a fixed 4-column offset.
    dose_cols: list[int] = []
    pos_col = 2  # column C holds the lateral position (cm); "/cm" header
    for row in ws.iter_rows(values_only=True):
        cells = list(row)
        for i, c in enumerate(cells):
            if isinstance(c, str) and c.strip().lower() == "dose ugy":
                dose_cols.append(i)
            if isinstance(c, str) and c.strip().lower() == "/cm":
                pos_col = i
        if dose_cols:
            break
    if not dose_cols:
        raise ValueError(
            "No 'Dose uGy' columns found in %s sheet %s" % (xlsx_path, sheet)
        )
    if mode_group < 0 or mode_group >= len(dose_cols):
        raise ValueError(
            "mode_group=%d out of range; %d modes available (0..%d)"
            % (mode_group, len(dose_cols), len(dose_cols) - 1)
        )
    dose_col = dose_cols[mode_group]
    hvl_col = dose_col + 2  # Dose, uGy/s, HVL -- consistent across modes

    # Second pass: read data rows (numeric lateral position in +/-20 cm).
    for row in ws.iter_rows(values_only=True):
        cells = list(row)
        if len(cells) <= hvl_col:
            continue
        c = cells[pos_col]
        if isinstance(c, (int, float)) and -20 < c < 20:
            pos.append(float(c))
            dose.append(
                float(cells[dose_col])
                if isinstance(cells[dose_col], (int, float))
                else float("nan")
            )
            hvl.append(
                float(cells[hvl_col])
                if isinstance(cells[hvl_col], (int, float))
                else float("nan")
            )
    wb.close()
    if not pos:
        raise ValueError("No profile rows found in %s sheet %s" % (xlsx_path, sheet))
    order = np.argsort(pos)
    pos_arr = np.asarray(pos)[order]
    return Profile(
        label="measured_group%d" % mode_group,
        position_cm=pos_arr,
        dose=np.asarray(dose)[order],
        hvl_mmAl=np.asarray(hvl)[order],
    )


def load_mc_profile(csv_path: str) -> Profile:
    """Load a scored MC cross-plane profile from a 2-column CSV (position_cm, dose)."""
    arr = np.loadtxt(csv_path, delimiter=",", comments="#", skiprows=1)
    if arr.ndim == 1:
        arr = arr.reshape(1, -1)
    order = np.argsort(arr[:, 0])
    return Profile(
        label="mc",
        position_cm=arr[order, 0],
        dose=arr[order, 1],
    )


def _normalise(values: np.ndarray) -> np.ndarray:
    peak = float(np.nanmax(values))
    if peak != peak or peak <= 0:  # NaN or non-positive -> cannot normalise
        return values
    return values / peak


def compare_profiles(
    measured: Profile,
    mc: Profile,
    output_csv: Optional[str] = None,
    output_png: Optional[str] = None,
) -> dict:
    """Compare measured and MC profiles (normalised to peak=1).

    The MC profile is interpolated onto the measured positions. Returns a dict
    with the per-position table and a peak-normalised RMS misfit.
    """
    mc_on_meas = np.interp(measured.position_cm, mc.position_cm, _normalise(mc.dose))
    meas_norm = _normalise(measured.dose)
    diff = mc_on_meas - meas_norm
    rms = float(np.sqrt(np.nanmean(diff**2))) if diff.size else float("nan")

    if output_csv:
        import csv

        with open(output_csv, "w", newline="") as f:
            w = csv.writer(f)
            w.writerow(["position_cm", "measured_norm", "mc_norm", "diff"])
            for p, m, c_, d in zip(measured.position_cm, meas_norm, mc_on_meas, diff):
                w.writerow(["%.3f" % p, "%.5f" % m, "%.5f" % c_, "%.5f" % d])
        logger.info("Wrote comparison CSV %s", output_csv)

    if output_png:
        _plot(measured, mc, mc_on_meas, output_png)

    return {
        "position_cm": measured.position_cm.tolist(),
        "measured_norm": meas_norm.tolist(),
        "mc_norm": mc_on_meas.tolist(),
        "rms_misfit": rms,
    }


def _plot(measured: Profile, mc: Profile, mc_interp: np.ndarray, path: str) -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(figsize=(7, 4))
    ax.plot(
        measured.position_cm, _normalise(measured.dose), "o-", label="measured", ms=3
    )
    ax.plot(measured.position_cm, mc_interp, "s-", label="MC (interp)", ms=3)
    ax.set_xlabel("Lateral position (cm)")
    ax.set_ylabel("Dose (peak-normalised)")
    ax.set_title("Bow-tie cross-plane profile: %s" % measured.label)
    ax.legend()
    ax.grid(True, alpha=0.3)
    fig.tight_layout()
    fig.savefig(path, dpi=120)
    plt.close(fig)
    logger.info("Wrote comparison PNG %s", path)
